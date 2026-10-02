"""Excel writer (openpyxl) with clickable links and urgency colours."""
from __future__ import annotations

from io import BytesIO
from typing import Optional

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from schemas import GapReport
from tracker import COLUMNS

_FILLS = {
    "Critical": "F8B4B4",
    "Soon": "FDE7A8",
    "OK": "BFE8C4",
    "Verify": "D9D9D9",
}
_WIDTHS = {
    "id": 6, "country": 16, "scholarship_name": 38, "provider": 28, "level": 10, "deadline": 13,
    "requirements": 30, "funding": 30, "official_link": 40, "confidence": 11, "fit_score": 9,
    "status": 11, "days_left": 10, "urgency": 10, "notes": 30,
}


def to_excel_bytes(df: pd.DataFrame, gap: Optional[GapReport] = None) -> bytes:
    out = BytesIO()
    sheet = df[[c for c in COLUMNS if c in df.columns]].copy()
    sheet["days_left"] = sheet["days_left"].astype(object).where(sheet["days_left"].notna(), None)

    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        sheet.to_excel(writer, sheet_name="Scholarships", index=False)
        ws = writer.sheets["Scholarships"]
        head_fill = PatternFill("solid", fgColor="1F3A5F")
        for idx, col in enumerate(sheet.columns, start=1):
            cell = ws.cell(row=1, column=idx)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = head_fill
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(idx)].width = _WIDTHS.get(col, 14)
        ws.freeze_panes = "C2"
        ws.auto_filter.ref = ws.dimensions

        cols = list(sheet.columns)
        link_i = cols.index("official_link") + 1 if "official_link" in cols else None
        urg_i = cols.index("urgency") + 1 if "urgency" in cols else None
        for r in range(2, ws.max_row + 1):
            for c in range(1, len(cols) + 1):
                ws.cell(row=r, column=c).alignment = Alignment(vertical="top", wrap_text=True)
            if link_i:
                cell = ws.cell(row=r, column=link_i)
                if cell.value and str(cell.value).startswith("http"):
                    cell.hyperlink = str(cell.value)
                    cell.font = Font(color="0563C1", underline="single")
            if urg_i:
                cell = ws.cell(row=r, column=urg_i)
                if cell.value in _FILLS:
                    cell.fill = PatternFill("solid", fgColor=_FILLS[cell.value])

        # Gap report sheet
        rows = []
        if gap:
            rows += [("Strength", s) for s in gap.strengths]
            rows += [("Gap", g) for g in gap.gaps]
            rows += [("Recommendation", r) for r in gap.recommendations]
        if "top_gaps" in df.columns:
            for _, r in df.iterrows():
                if r.get("top_gaps"):
                    rows.append((f"Gaps for: {r['scholarship_name']}", r["top_gaps"]))
        gdf = pd.DataFrame(rows, columns=["Type", "Detail"]) if rows else pd.DataFrame({"Type": [], "Detail": []})
        gdf.to_excel(writer, sheet_name="Gap Report", index=False)
        gs = writer.sheets["Gap Report"]
        gs.column_dimensions["A"].width = 44
        gs.column_dimensions["B"].width = 100
        for c in (1, 2):
            gs.cell(row=1, column=c).font = Font(bold=True)
        for r in range(2, gs.max_row + 1):
            for c in (1, 2):
                gs.cell(row=r, column=c).alignment = Alignment(vertical="top", wrap_text=True)
    return out.getvalue()
